#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
planilhas.py — Gerador de PLANILHAS por template (determinístico, openpyxl).

Diferente da exportação (que espelha uma tabela markdown), aqui o Python MONTA
uma planilha .xlsx real, completa e sempre igual: abas, formatação de moeda,
fórmulas (SUMIFS por mês, saldo acumulado), validação de dados (dropdowns) e
orçamento. Não depende do modelo de linguagem.

API:
    listar() -> [ {id, nome, descricao, campos:[{id,label,tipo,default,ajuda}]} ]
    gerar(template_id, params) -> (bytes, filename)

Fórmulas: escritas em INGLÊS com vírgula (padrão do arquivo .xlsx); o Excel em
pt-BR exibe traduzido (SOMASES etc.) automaticamente.
"""
from __future__ import annotations

import datetime as _dt
import io
import re
from typing import List, Tuple

MESES_PT = ["jan", "fev", "mar", "abr", "mai", "jun",
            "jul", "ago", "set", "out", "nov", "dez"]

_CUR = 'R$ #,##0.00'


# ---------------------------------------------------------------------------
# Catálogo de templates
# ---------------------------------------------------------------------------

def _ano_atual() -> int:
    return _dt.date.today().year


def listar() -> List[dict]:
    ano = _ano_atual()
    return [
        {
            "id": "controle_financeiro_pf",
            "nome": "Controle Financeiro — Pessoa Física",
            "descricao": "Lançamentos + Resumo mensal (receitas/despesas/saldo acumulado) "
                         "+ categorias com dropdown + orçamento por categoria. Com fórmulas e moeda.",
            "campos": [
                {"id": "ano", "label": "Ano", "tipo": "number", "default": ano,
                 "ajuda": "Ano de referência do resumo mensal."},
                {"id": "saldo_inicial", "label": "Saldo inicial (R$)", "tipo": "number", "default": 0,
                 "ajuda": "Saldo em caixa no começo do ano."},
                {"id": "categorias_receita", "label": "Categorias de receita (uma por linha)",
                 "tipo": "textarea",
                 "default": "Salário\nFreelas / Extra\nRendimentos\nReembolsos\nOutras receitas",
                 "ajuda": "Aparecem no dropdown de Categoria."},
                {"id": "categorias_despesa", "label": "Categorias de despesa (uma por linha)",
                 "tipo": "textarea",
                 "default": ("Moradia\nAlimentação\nTransporte\nSaúde\nEducação\nLazer\n"
                             "Contas (água/luz/internet)\nAssinaturas\nDívidas / Empréstimos\n"
                             "Investimentos\nReserva de emergência\nOutras despesas"),
                 "ajuda": "Aparecem no dropdown e no orçamento."},
            ],
        },
    ]


def gerar(template_id: str, params: dict) -> Tuple[bytes, str]:
    if template_id == "controle_financeiro_pf":
        return _controle_financeiro_pf(params or {})
    raise ValueError(f"Template desconhecido: {template_id}")


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _linhas(txt: str, fallback: List[str]) -> List[str]:
    itens = [l.strip() for l in str(txt or "").replace("\r", "").split("\n")]
    itens = [l for l in itens if l]
    # remove duplicatas preservando ordem
    vistos, out = set(), []
    for l in itens:
        k = l.lower()
        if k not in vistos:
            vistos.add(k); out.append(l)
    return out or list(fallback)


def _int(v, default):
    try:
        return int(float(str(v).replace(",", ".")))
    except (TypeError, ValueError):
        return default


def _float(v, default=0.0):
    try:
        return float(str(v).replace(".", "").replace(",", ".")) if isinstance(v, str) and "," in v \
            else float(v)
    except (TypeError, ValueError):
        return default


# ---------------------------------------------------------------------------
# Template: Controle Financeiro — Pessoa Física
# ---------------------------------------------------------------------------

def _controle_financeiro_pf(p: dict) -> Tuple[bytes, str]:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.utils import get_column_letter

    ano = _int(p.get("ano"), _ano_atual())
    saldo_inicial = _float(p.get("saldo_inicial"), 0.0)
    cat_rec = _linhas(p.get("categorias_receita"),
                      ["Salário", "Rendimentos", "Outras receitas"])
    cat_desp = _linhas(p.get("categorias_despesa"),
                       ["Moradia", "Alimentação", "Transporte", "Outras despesas"])
    tipos = ["Receita", "Despesa", "Investimento"]

    # estilos
    AZUL = "1F3B57"; CINZA = "E9EEF3"; AMBAR = "FFF3E0"
    h_font = Font(bold=True, color="FFFFFF", size=11)
    h_fill = PatternFill("solid", fgColor=AZUL)
    sub_fill = PatternFill("solid", fgColor=CINZA)
    tot_fill = PatternFill("solid", fgColor=AMBAR)
    center = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left", vertical="center")
    thin = Side(style="thin", color="B7C2CC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    def header(ws, row, titulos, larguras=None):
        for i, t in enumerate(titulos, 1):
            c = ws.cell(row=row, column=i, value=t)
            c.font = h_font; c.fill = h_fill; c.alignment = center; c.border = border
        if larguras:
            for i, w in enumerate(larguras, 1):
                ws.column_dimensions[get_column_letter(i)].width = w

    wb = Workbook()
    N = 300  # linhas de lançamento pré-formatadas

    # ---- Aba LANÇAMENTOS ----
    ws = wb.active
    ws.title = "Lançamentos"
    ws.sheet_view.showGridLines = True
    ws["A1"] = f"LANÇAMENTOS — {ano}"
    ws["A1"].font = Font(bold=True, size=13, color=AZUL)
    ws.merge_cells("A1:F1")
    header(ws, 3, ["Data", "Tipo", "Categoria", "Descrição", "Valor", "Mês"],
           [13, 15, 26, 40, 15, 12])
    ws.freeze_panes = "A4"
    for r in range(4, 4 + N):
        ws.cell(row=r, column=1).number_format = "dd/mm/yyyy"
        ws.cell(row=r, column=5).number_format = _CUR
        # coluna Mês (auto): mostra o mês da data p/ conferência
        ws.cell(row=r, column=6,
                value=f'=IF($A{r}="","",TEXT($A{r},"mmm/yyyy"))')
        ws.cell(row=r, column=6).alignment = center
        for col in range(1, 7):
            ws.cell(row=r, column=col).border = border
    # validações (dropdowns)
    dv_tipo = DataValidation(type="list", formula1='"%s"' % ",".join(tipos), allow_blank=True)
    dv_tipo.error = "Escolha: " + ", ".join(tipos)
    dv_tipo.prompt = "Receita, Despesa ou Investimento"
    ws.add_data_validation(dv_tipo); dv_tipo.add(f"B4:B{3+N}")
    ncat = len(cat_rec) + len(cat_desp)
    dv_cat = DataValidation(type="list",
                            formula1=f"=Categorias!$A$2:$A${1+ncat}", allow_blank=True)
    ws.add_data_validation(dv_cat); dv_cat.add(f"C4:C{3+N}")
    # linha de total rápida no topo
    ws["H3"] = "Entradas"; ws["H4"] = "Saídas"; ws["H5"] = "Saldo"
    for cell in ("H3", "H4", "H5"):
        ws[cell].font = Font(bold=True)
    ws["I3"] = f'=SUMIFS($E$4:$E${3+N},$B$4:$B${3+N},"Receita")'
    ws["I4"] = f'=SUMIFS($E$4:$E${3+N},$B$4:$B${3+N},"Despesa")'
    ws["I5"] = "=I3-I4"
    for cell in ("I3", "I4", "I5"):
        ws[cell].number_format = _CUR; ws[cell].font = Font(bold=True)
    ws.column_dimensions["H"].width = 12; ws.column_dimensions["I"].width = 15

    LAN = "'Lançamentos'"

    # ---- Aba RESUMO ----
    rs = wb.create_sheet("Resumo")
    rs["A1"] = f"RESUMO MENSAL — {ano}"
    rs["A1"].font = Font(bold=True, size=13, color=AZUL); rs.merge_cells("A1:E1")
    rs["A2"] = "Saldo inicial"; rs["A2"].font = Font(bold=True)
    rs["B2"] = saldo_inicial; rs["B2"].number_format = _CUR
    header(rs, 4, ["Mês", "Receitas", "Despesas", "Saldo do mês", "Saldo acumulado"],
           [14, 16, 16, 16, 18])
    rs.freeze_panes = "A5"
    first_row = 5
    for m in range(12):
        r = first_row + m
        d = _dt.date(ano, m + 1, 1)
        cA = rs.cell(row=r, column=1, value=d); cA.number_format = "mmm/yyyy"; cA.alignment = center
        # Receitas / Despesas do mês via SUMIFS por faixa de datas
        rng = (f'{LAN}!$E$4:$E${3+N}')
        colB = f'{LAN}!$B$4:$B${3+N}'
        colA = f'{LAN}!$A$4:$A${3+N}'
        rs.cell(row=r, column=2,
                value=f'=SUMIFS({rng},{colB},"Receita",{colA},">="&$A{r},{colA},"<="&EOMONTH($A{r},0))')
        rs.cell(row=r, column=3,
                value=f'=SUMIFS({rng},{colB},"Despesa",{colA},">="&$A{r},{colA},"<="&EOMONTH($A{r},0))')
        rs.cell(row=r, column=4, value=f'=B{r}-C{r}')
        if m == 0:
            rs.cell(row=r, column=5, value=f'=$B$2+D{r}')
        else:
            rs.cell(row=r, column=5, value=f'=E{r-1}+D{r}')
        for col in range(2, 6):
            cc = rs.cell(row=r, column=col); cc.number_format = _CUR; cc.border = border
        rs.cell(row=r, column=1).border = border
    # total
    tr = first_row + 12
    rs.cell(row=tr, column=1, value="TOTAL").font = Font(bold=True)
    rs.cell(row=tr, column=2, value=f"=SUM(B{first_row}:B{tr-1})")
    rs.cell(row=tr, column=3, value=f"=SUM(C{first_row}:C{tr-1})")
    rs.cell(row=tr, column=4, value=f"=B{tr}-C{tr}")
    rs.cell(row=tr, column=5, value=f"=E{tr-1}")
    for col in range(1, 6):
        cc = rs.cell(row=tr, column=col); cc.fill = tot_fill; cc.font = Font(bold=True); cc.border = border
        if col >= 2:
            cc.number_format = _CUR

    # ---- Aba CATEGORIAS (fonte do dropdown + gasto por categoria) ----
    cs = wb.create_sheet("Categorias")
    cs["A1"] = "Categorias"; cs["A1"].font = h_font; cs["A1"].fill = h_fill; cs["A1"].alignment = center
    cs.column_dimensions["A"].width = 28
    todas = cat_rec + cat_desp
    for i, cat in enumerate(todas, start=2):
        cc = cs.cell(row=i, column=1, value=cat)
        cc.fill = sub_fill if cat in cat_rec else PatternFill("solid", fgColor="FBE9E7")
        cc.border = border
    # gasto por categoria de despesa
    header(cs, 1, ["Categorias"], [28])
    cs["C1"] = "Despesa por categoria"; cs["C1"].font = h_font; cs["C1"].fill = h_fill; cs["C1"].alignment = center
    cs["D1"] = "Total"; cs["D1"].font = h_font; cs["D1"].fill = h_fill; cs["D1"].alignment = center
    cs.column_dimensions["C"].width = 26; cs.column_dimensions["D"].width = 15
    for i, cat in enumerate(cat_desp, start=2):
        cs.cell(row=i, column=3, value=cat).border = border
        cc = cs.cell(row=i, column=4,
                     value=f'=SUMIFS({LAN}!$E$4:$E${3+N},{LAN}!$C$4:$C${3+N},C{i},{LAN}!$B$4:$B${3+N},"Despesa")')
        cc.number_format = _CUR; cc.border = border

    # ---- Aba ORÇAMENTO (planejado x real) ----
    os_ = wb.create_sheet("Orçamento")
    os_["A1"] = f"ORÇAMENTO MENSAL — despesas — {ano}"
    os_["A1"].font = Font(bold=True, size=12, color=AZUL); os_.merge_cells("A1:D1")
    header(os_, 3, ["Categoria", "Planejado (mês)", "Real (ano)", "Diferença (ano)"],
           [26, 16, 16, 16])
    os_.freeze_panes = "A4"
    for i, cat in enumerate(cat_desp, start=4):
        os_.cell(row=i, column=1, value=cat).border = border
        cp = os_.cell(row=i, column=2, value=0); cp.number_format = _CUR; cp.border = border  # usuário preenche
        cr = os_.cell(row=i, column=3,
                      value=f'=SUMIFS({LAN}!$E$4:$E${3+N},{LAN}!$C$4:$C${3+N},A{i},{LAN}!$B$4:$B${3+N},"Despesa")')
        cr.number_format = _CUR; cr.border = border
        cd = os_.cell(row=i, column=4, value=f'=B{i}*12-C{i}')
        cd.number_format = _CUR; cd.border = border

    # ---- Aba INSTRUÇÕES ----
    ins = wb.create_sheet("Leia-me")
    ins.column_dimensions["A"].width = 100
    linhas = [
        "CONTROLE FINANCEIRO — PESSOA FÍSICA",
        "",
        "Como usar:",
        "1) Aba LANÇAMENTOS: registre cada movimento — Data, Tipo (dropdown), Categoria (dropdown),",
        "   Descrição e Valor. O restante é calculado sozinho.",
        "2) Aba RESUMO: receitas, despesas, saldo do mês e saldo acumulado, por mês. Ajuste o Saldo inicial (B2).",
        "3) Aba CATEGORIAS: edite a lista (col. A) para mudar os dropdowns. Veja o gasto por categoria.",
        "4) Aba ORÇAMENTO: preencha o Planejado por categoria; o Real e a Diferença vêm das fórmulas.",
        "",
        "Boas práticas embutidas: categorização, moeda R$, resumo mensal, saldo acumulado, orçamento,",
        "reserva de emergência (categoria) e separação receita/despesa/investimento.",
        "",
        "Gerado pelo NÚCLEO IA PORTÁTIL — offline, sem IA (planilha determinística).",
    ]
    for i, t in enumerate(linhas, start=1):
        c = ins.cell(row=i, column=1, value=t)
        if i == 1:
            c.font = Font(bold=True, size=13, color=AZUL)

    # ordem das abas
    wb.move_sheet("Leia-me", -(len(wb.sheetnames) - 1))

    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue(), f"Controle_Financeiro_PF_{ano}.xlsx"
