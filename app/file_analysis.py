#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
NÚCLEO IA PORTÁTIL — Análise determinística de arquivos.

V0.6.5.2 — Leitura Semântica de Planilhas
------------------------------------------
Além da leitura tabular tradicional (cabeçalho + linhas), o extrator agora
classifica cada aba de XLSX como:

    TABELA               -> linhas / colunas  (comportamento V0.6.5.x)
    FORMULÁRIO/CALCULADORA -> seções / pares rótulo->valor

Para formulários/calculadoras o extrator:
  - detecta seções (linhas só-texto que titulam blocos);
  - associa rótulo -> valor por adjacência, preservando a coordenada da célula;
  - diferencia vazio estrutural de dado realmente ausente;
  - interpreta formatação numérica (moeda / percentual / inteiro / decimal /
    data / hora) usando xl/styles.xml;
  - associa fórmulas ao rótulo mais próximo.

Segurança preservada: nenhuma macro/código é executado; XLSX é tratado como
ZIP/XML com limites de expansão; fórmulas são LIDAS, nunca recalculadas.
"""
from __future__ import annotations

import csv
import io
import json
import math
import re
import statistics
import zipfile
from collections import Counter, defaultdict
from datetime import date, timedelta
from pathlib import Path
import xml.etree.ElementTree as ET

MAX_FILE_BYTES = 50 * 1024 * 1024
MAX_ROWS = 50_000
MAX_COLS = 128
MAX_SHEETS = 12
MAX_JSON_DEPTH = 6
MAX_XLSX_UNCOMPRESSED = 250 * 1024 * 1024
MAX_XLSX_MEMBERS = 2500
MAX_XLSX_MEMBER = 64 * 1024 * 1024
MAX_FORM_ITEMS = 200
MAX_PDF_PAGES = 800          # páginas processadas no máximo
MAX_PDF_TEXT_CHARS = 200_000  # teto do texto extraído acumulado
SUPPORTED_EXTENSIONS = {".xlsx", ".csv", ".txt", ".md", ".json", ".pdf", ".docx"}

_CELL_REF_RE = re.compile(r"(\$?[A-Z]{1,3})(\$?\d+)")
_NUM_RE = re.compile(r"^[+-]?(?:\d+(?:[.,]\d+)?|\d*[.,]\d+)(?:[eE][+-]?\d+)?$")
_XLNS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"

# numFmtId embutidos do OOXML (subconjunto relevante).
_BUILTIN_FMT = {
    0: None, 1: "integer", 2: "decimal", 3: "integer", 4: "decimal",
    9: "percent", 10: "percent",
    14: "date", 15: "date", 16: "date", 17: "date",
    18: "time", 19: "time", 20: "time", 21: "time", 22: "date",
    37: "integer", 38: "integer", 39: "decimal", 40: "decimal",
    41: "currency", 42: "currency", 43: "currency", 44: "currency",
    45: "time", 46: "time", 47: "time", 48: "decimal", 49: None,
}


class FileAnalysisError(RuntimeError):
    pass


def _decode_text(data: bytes) -> tuple[str, str]:
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return data.decode(enc), enc
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace"), "utf-8-replace"


def _clean_scalar(value):
    if value is None:
        return None
    if isinstance(value, (int, float, bool)):
        return value
    text = str(value).strip()
    return text if text else None


def _to_number(value):
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        x = float(value)
        return x if math.isfinite(x) else None
    s = str(value).strip().replace(" ", " ")
    if not s:
        return None
    # Conservative locale handling.
    s = s.replace("R$", "").replace("%", "").strip()
    if not _NUM_RE.match(s.replace(" ", "")):
        return None
    if "," in s and "." in s:
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        x = float(s)
        return x if math.isfinite(x) else None
    except ValueError:
        return None


def _norm_formula(formula: str) -> str:
    s = str(formula or "").strip().upper()
    s = _CELL_REF_RE.sub(lambda m: m.group(1) + "#", s)
    s = re.sub(r"\s+", "", s)
    return s[:300]


# ---------------------------------------------------------------------------
# Formatação numérica (V0.6.5.2)
# ---------------------------------------------------------------------------
def _br_number(n: float, decimals: int | None = None) -> str:
    neg = n < 0
    n = abs(n)
    if decimals is None:
        if float(n).is_integer():
            s = f"{int(n):,}".replace(",", ".")
            return ("-" if neg else "") + s
        decimals = 2
    s = f"{n:,.{decimals}f}".replace(",", "\x00").replace(".", ",").replace("\x00", ".")
    return ("-" if neg else "") + s


def _excel_date(serial: float) -> str | None:
    try:
        if serial < 1 or serial > 60_000:
            return None
        return (date(1899, 12, 30) + timedelta(days=int(serial))).isoformat()
    except Exception:
        return None


def _format_cell_value(value, kind):
    if value is None:
        return None
    n = _to_number(value)
    if n is None:
        return str(value)[:160]
    if kind == "percent":
        p = n * 100
        s = f"{p:.2f}".rstrip("0").rstrip(".")
        return f"{s}%"
    if kind == "currency":
        return "R$ " + _br_number(n, 2)
    if kind == "integer":
        return str(int(round(n)))
    if kind == "date":
        return _excel_date(n) or _br_number(n)
    if kind in ("time", "hours"):
        return _br_number(n)
    if float(n).is_integer():
        return str(int(n))
    return _br_number(n)


def _fmt_kind_from_code(code) -> str | None:
    c = str(code or "").lower()
    if not c or c == "general":
        return None
    if "%" in c:
        return "percent"
    if any(sym in c for sym in ("r$", "$", "€", "£")):
        return "currency"
    if "[h" in c or ":mm" in c or "hh" in c:
        return "time"
    if any(x in c for x in ("yy", "dd", "mmm")) or ("m" in c and "y" in c):
        return "date"
    if "0.0" in c or "#,##0.0" in c:
        return "decimal"
    if "0" in c or "#" in c:
        return "integer"
    return None


def _xlsx_number_formats(z) -> dict:
    """style_index (cellXfs) -> kind ('currency'|'percent'|'date'|...|None)."""
    name = "xl/styles.xml"
    if name not in z.namelist():
        return {}
    try:
        root = ET.fromstring(z.read(name))
    except Exception:
        return {}
    custom = {}
    numfmts = root.find(_XLNS + "numFmts")
    if numfmts is not None:
        for nf in numfmts:
            fid = nf.attrib.get("numFmtId")
            code = nf.attrib.get("formatCode")
            if fid is not None:
                try:
                    custom[int(fid)] = _fmt_kind_from_code(code)
                except ValueError:
                    pass
    kinds = {}
    cellxfs = root.find(_XLNS + "cellXfs")
    if cellxfs is not None:
        for i, xf in enumerate(cellxfs):
            try:
                fid = int(xf.attrib.get("numFmtId", "0"))
            except ValueError:
                fid = 0
            if fid in custom:
                kinds[i] = custom[fid]
            else:
                kinds[i] = _BUILTIN_FMT.get(fid)
    return kinds


class ColumnStats:
    def __init__(self, name):
        self.name = str(name)
        self.count = 0
        self.missing = 0
        self.numeric_count = 0
        self.sum = 0.0
        self.min = None
        self.max = None
        self.values = Counter()
        self.samples = []

    def add(self, value):
        self.count += 1
        v = _clean_scalar(value)
        if v is None:
            self.missing += 1
            return
        if len(self.samples) < 6:
            self.samples.append(str(v)[:120])
        n = _to_number(v)
        if n is not None:
            self.numeric_count += 1
            self.sum += n
            self.min = n if self.min is None else min(self.min, n)
            self.max = n if self.max is None else max(self.max, n)
        if len(self.values) < 5000 or str(v) in self.values:
            self.values[str(v)[:160]] += 1

    def summary(self):
        non_missing = self.count - self.missing
        numeric_ratio = (self.numeric_count / non_missing) if non_missing else 0
        out = {
            "name": self.name,
            "count": self.count,
            "missing": self.missing,
            "numeric": numeric_ratio >= 0.8 and self.numeric_count > 0,
            "samples": self.samples,
        }
        if out["numeric"]:
            out.update({
                "numeric_count": self.numeric_count,
                "sum": round(self.sum, 6),
                "mean": round(self.sum / self.numeric_count, 6) if self.numeric_count else None,
                "min": self.min,
                "max": self.max,
            })
        else:
            out["unique_approx"] = len(self.values)
            out["top_values"] = [
                {"value": k, "count": v}
                for k, v in self.values.most_common(8)
            ]
        return out


def _analyze_rows(rows, source_name="dados", max_rows=MAX_ROWS):
    rows = iter(rows)
    try:
        first = next(rows)
    except StopIteration:
        return {
            "kind": "table",
            "name": source_name,
            "row_count": 0,
            "columns": [],
            "sample_rows": [],
            "groups": [],
        }

    first = list(first)[:MAX_COLS]
    headers = []
    seen = set()
    for i, value in enumerate(first):
        name = str(value).strip() if value is not None else ""
        if not name:
            name = f"coluna_{i+1}"
        base = name[:80]
        candidate = base
        k = 2
        while candidate.lower() in seen:
            candidate = f"{base}_{k}"
            k += 1
        seen.add(candidate.lower())
        headers.append(candidate)

    stats = [ColumnStats(h) for h in headers]
    sample_rows = []
    row_count = 0
    group_raw = defaultdict(lambda: defaultdict(lambda: [0.0, 0]))

    def handle(row):
        nonlocal row_count
        if row_count >= max_rows:
            return False
        vals = list(row)[:len(headers)]
        vals += [None] * (len(headers) - len(vals))
        row_count += 1
        for i, v in enumerate(vals):
            stats[i].add(v)
        if len(sample_rows) < 12:
            sample_rows.append({
                headers[i]: _clean_scalar(vals[i])
                for i in range(len(headers))
            })
        return True

    for row in rows:
        if not handle(row):
            break

    summaries = [s.summary() for s in stats]
    # V0.6.5.1: colunas totalmente vazias só poluem análises de planilhas
    # tipo formulário. Mantemos apenas colunas com ao menos um valor útil.
    summaries = [
        s for s in summaries
        if int(s.get("missing") or 0) < int(s.get("count") or 0)
    ]
    numeric_idxs = [i for i, s in enumerate(summaries) if s.get("numeric")]
    categorical_idxs = [
        i for i, s in enumerate(summaries)
        if not s.get("numeric")
        and 2 <= int(s.get("unique_approx") or 0) <= 16
    ][:5]

    # Rebuild group summaries from samples is too weak; keep deterministic global stats
    # and categorical distributions. Query context will combine both safely.
    groups = []
    for i in categorical_idxs:
        s = summaries[i]
        groups.append({
            "column": s.get("name"),
            "values": s.get("top_values", []),
        })

    return {
        "kind": "table",
        "name": source_name,
        "row_count": row_count,
        "truncated": row_count >= max_rows,
        "columns": summaries,
        "sample_rows": sample_rows,
        "groups": groups,
    }


def analyze_csv(path: Path):
    data = path.read_bytes()
    text, encoding = _decode_text(data)
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
        dialect.delimiter = ";"
    reader = csv.reader(io.StringIO(text), dialect)
    result = _analyze_rows(reader, source_name=path.name)
    result.update({
        "type": "csv",
        "encoding": encoding,
        "delimiter": dialect.delimiter,
        "structure": "table",
    })
    return result


def _safe_xlsx_zip(path: Path):
    z = zipfile.ZipFile(path, "r")
    infos = z.infolist()
    if len(infos) > MAX_XLSX_MEMBERS:
        z.close()
        raise FileAnalysisError("XLSX possui arquivos internos demais.")
    total = 0
    for info in infos:
        total += info.file_size
        if info.file_size > MAX_XLSX_MEMBER:
            z.close()
            raise FileAnalysisError("XLSX possui um componente interno grande demais.")
    if total > MAX_XLSX_UNCOMPRESSED:
        z.close()
        raise FileAnalysisError("XLSX expandido excede o limite de segurança.")
    return z


def _xlsx_shared_strings(z):
    name = "xl/sharedStrings.xml"
    if name not in z.namelist():
        return []
    out = []
    with z.open(name) as f:
        root = ET.parse(f).getroot()
    for si in root:
        parts = []
        for elem in si.iter():
            if elem.tag.endswith("}t") and elem.text:
                parts.append(elem.text)
        out.append("".join(parts))
    return out


def _xlsx_sheet_map(z):
    ns_main = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    wb = ET.fromstring(z.read("xl/workbook.xml"))
    relroot = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
    rels = {}
    for rel in relroot:
        rid = rel.attrib.get("Id")
        target = rel.attrib.get("Target")
        if rid and target:
            if target.startswith("/"):
                full = target.lstrip("/")
            else:
                full = "xl/" + target.lstrip("/")
            rels[rid] = full.replace("\\", "/")
    sheets = []
    for sheet in wb.findall(".//m:sheets/m:sheet", ns_main):
        name = sheet.attrib.get("name", "Planilha")
        rid = sheet.attrib.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
        target = rels.get(rid)
        if target and target in z.namelist():
            sheets.append((name, target))
    return sheets[:MAX_SHEETS]


def _xlsx_cell_value(cell, shared):
    t = cell.attrib.get("t")
    v = None
    formula = None
    for child in cell:
        if child.tag.endswith("}f"):
            formula = child.text or ""
        elif child.tag.endswith("}v"):
            v = child.text
        elif child.tag.endswith("}is"):
            parts = []
            for elem in child.iter():
                if elem.tag.endswith("}t") and elem.text:
                    parts.append(elem.text)
            v = "".join(parts)
    if t == "s" and v is not None:
        try:
            idx = int(v)
            v = shared[idx] if 0 <= idx < len(shared) else v
        except ValueError:
            pass
    elif t == "b":
        v = True if v == "1" else False
    elif t in ("n", None) and v is not None:
        try:
            x = float(v)
            v = int(x) if x.is_integer() else x
        except ValueError:
            pass
    return v, formula


def _col_index(cell_ref: str) -> int:
    m = re.match(r"([A-Z]+)", str(cell_ref or "").upper())
    if not m:
        return 0
    n = 0
    for ch in m.group(1):
        n = n * 26 + (ord(ch) - 64)
    return max(0, n - 1)


def _xlsx_grid(z, target, shared, style_kinds):
    """Lê a aba inteira como grade esparsa: (row, col) -> célula rica.

    Retorna (grid, max_row, max_col, formulas). Base tanto para o ramo
    tabular quanto para o ramo semântico (formulário).
    """
    grid = {}
    formulas = []
    max_row = 0
    max_col = 0
    with z.open(target) as f:
        context = ET.iterparse(f, events=("end",))
        for _, elem in context:
            if not elem.tag.endswith("}row"):
                continue
            try:
                r = int(elem.attrib.get("r") or 0)
            except ValueError:
                r = 0
            if r <= 0 or r > MAX_ROWS:
                elem.clear()
                continue
            for cell in list(elem):
                if not cell.tag.endswith("}c"):
                    continue
                ref = cell.attrib.get("r", "")
                c = _col_index(ref)
                if c >= MAX_COLS:
                    continue
                value, formula = _xlsx_cell_value(cell, shared)
                if value is None and not formula:
                    continue
                s_idx = cell.attrib.get("s")
                kind = None
                if s_idx and s_idx.isdigit():
                    kind = style_kinds.get(int(s_idx))
                grid[(r, c)] = {
                    "ref": ref, "row": r, "col": c,
                    "value": value, "formula": formula, "kind": kind,
                }
                if formula:
                    formulas.append({
                        "cell": ref,
                        "formula": formula[:500],
                        "pattern": _norm_formula(formula),
                        "column": re.sub(r"\d+", "", ref),
                    })
                if r > max_row:
                    max_row = r
                if c > max_col:
                    max_col = c
            elem.clear()
    if len(formulas) > 20_000:
        formulas[:] = formulas[:20_000]
    return grid, max_row, max_col, formulas


def _grid_to_rows(grid, max_row, max_col):
    """Materializa a grade em linhas (para o ramo tabular). Descarta linhas
    totalmente vazias no topo para que a primeira linha real vire cabeçalho."""
    rows = []
    started = False
    for r in range(1, max_row + 1):
        row = [None] * (max_col + 1)
        has = False
        for c in range(max_col + 1):
            cell = grid.get((r, c))
            if cell is not None:
                row[c] = cell["value"]
                has = True
        if not started and not has:
            continue
        started = True
        rows.append(row)
    return rows


def _cell_is_value(cell):
    return bool(cell.get("formula")) or _to_number(cell.get("value")) is not None


def _cell_is_text(cell):
    if _cell_is_value(cell):
        return False
    v = cell.get("value")
    return v is not None and str(v).strip() != ""


def _classify_structure(grid, max_row, max_col):
    """Heurística TABELA vs FORMULÁRIO por ADJACÊNCIA rótulo->valor.

    Uma calculadora/formulário — inclusive em blocos lado a lado — tem quase
    todo valor precedido, na mesma linha, por uma célula de texto (o rótulo).
    Uma tabela de dados tem os valores em grade sob um cabeçalho, então a
    maioria dos valores é precedida por outro valor, não por texto.

    Salvaguarda: se a primeira linha preenchida for um cabeçalho textual com
    >= 3 colunas e houver >= 3 linhas de dados sob ele, tratamos como TABELA.
    """
    if not grid:
        return "table"

    rows = defaultdict(dict)
    for (r, c), cell in grid.items():
        rows[r][c] = cell

    # Salvaguarda de tabela com cabeçalho textual largo.
    first_r = min(rows)
    first = rows[first_r]
    header_cols = sorted(first)
    if len(header_cols) >= 3 and all(_cell_is_text(first[c]) for c in header_cols):
        data_rows = 0
        for r in sorted(rows):
            if r == first_r:
                continue
            vals = sum(1 for c in header_cols
                       if c in rows[r] and _cell_is_value(rows[r][c]))
            if vals >= 2:
                data_rows += 1
        if data_rows >= 3:
            return "table"

    # Razão de valores rotulados (texto imediatamente à esquerda).
    total_vals = 0
    labeled = 0
    for (r, c), cell in grid.items():
        if _cell_is_value(cell):
            total_vals += 1
            left = grid.get((r, c - 1))
            if left is not None and _cell_is_text(left):
                labeled += 1
    if total_vals >= 3 and (labeled / total_vals) >= 0.6:
        return "form"
    return "table"


def _column_blocks(grid, max_col):
    """Agrupa colunas usadas em blocos contíguos (separados por coluna vazia).

    Permite tratar planilhas com dois formulários lado a lado (ex.: B/C e E/F)
    como blocos independentes, preservando as seções de cada um.
    """
    used = sorted({c for (_, c) in grid})
    if not used:
        return []
    blocks = []
    start = prev = used[0]
    for c in used[1:]:
        if c == prev + 1:
            prev = c
            continue
        blocks.append((start, prev))
        start = prev = c
    blocks.append((start, prev))
    return blocks


def _looks_like_section(text) -> bool:
    t = str(text or "").strip()
    if len(t) < 4:
        return False
    if re.match(r"(?i)^se[cç][aã]o\b", t):
        return True
    letters = re.sub(r"[^A-Za-zÀ-ÿ]", "", t)
    if letters and sum(1 for ch in letters if ch.isupper()) / len(letters) >= 0.7:
        return True
    return False


_MONEY_LABEL_RE = re.compile(
    r"(?i)\b(valor|custo|custos|pre[cç]o|sal[aá]rio|or[cç]amento|"
    r"receita|lucro|faturamento|subtotal|total de custos|r\$)\b"
)
_NONMONEY_LABEL_RE = re.compile(
    r"(?i)\b(hora|horas|dia|dias|semana|semanas|m[eê]s(?:es)?|"
    r"quantidade|qtd|n[uú]mero|percentual|margem|%|estimativa)\b"
)


def _effective_kind(label, kind):
    """Evita rotular como moeda (R$) valores cujo rótulo indica outra unidade
    (horas, dias, percentual...). O estilo de coluna do Excel costuma marcar
    a coluna inteira como moeda, mesmo em células que não são dinheiro.

    Rótulos com palavra monetária (valor, custo, preço, salário...) têm
    precedência e permanecem como moeda mesmo contendo 'hora'."""
    if kind != "currency" or not label:
        return kind
    s = str(label)
    if _MONEY_LABEL_RE.search(s):
        return "currency"
    if _NONMONEY_LABEL_RE.search(s):
        return None
    return kind


def _extract_form_block(grid, cols, max_items):
    """Extrai seções + pares rótulo->valor de UM bloco de colunas (cols=range)."""
    colset = set(cols)
    by_row = defaultdict(list)
    for (r, c), cell in grid.items():
        if c in colset:
            by_row[r].append((c, cell))

    sections = []
    current = {"title": None, "items": []}
    total = 0
    pending_label = None  # rótulo cujo valor pode estar na linha de baixo

    def flush():
        if current["title"] is not None or current["items"]:
            sections.append({"title": current["title"], "items": list(current["items"])})

    for r in sorted(by_row):
        if total >= max_items:
            break
        cells = sorted(by_row[r], key=lambda x: x[0])
        has_value = any(_cell_is_value(cell) for _, cell in cells)
        has_text = any(_cell_is_text(cell) for _, cell in cells)

        # Linha só-texto: seção (se parecer título) ou rótulo pendente.
        if has_text and not has_value:
            label = " ".join(str(cell["value"]).strip()
                             for _, cell in cells if _cell_is_text(cell))
            if _looks_like_section(label):
                flush()
                current = {"title": label[:140], "items": []}
                pending_label = None
            else:
                pending_label = label
            continue

        # Linha com valores: pareia cada valor com o texto imediatamente à
        # esquerda (na mesma linha); se não houver, usa o rótulo pendente.
        last_label = None
        for c, cell in cells:
            if _cell_is_text(cell):
                last_label = str(cell["value"]).strip()
                continue
            if not _cell_is_value(cell):
                continue
            lbl = last_label or pending_label or "(sem rótulo)"
            val = cell.get("value")
            kind = _effective_kind(lbl, cell.get("kind"))
            current["items"].append({
                "label": str(lbl)[:140],
                "value": val,
                "formatted": _format_cell_value(val, kind),
                "cell": cell["ref"],
                "formula": cell.get("formula"),
                "kind": kind,
            })
            total += 1
            last_label = None
            pending_label = None
            if total >= max_items:
                break

    flush()
    return sections


def _extract_form(grid, max_row, max_col, max_items=MAX_FORM_ITEMS):
    """Extrai seções e pares rótulo->valor, ciente de blocos lado a lado.

    Cada bloco de colunas contíguas (separado por coluna vazia) é extraído de
    cima para baixo de forma independente, preservando as seções de cada bloco.
    A ordem final é da esquerda para a direita (bloco a bloco).
    """
    blocks = _column_blocks(grid, max_col)
    if not blocks:
        return []
    sections = []
    remaining = max_items
    for (c0, c1) in blocks:
        if remaining <= 0:
            break
        block_cols = range(c0, c1 + 1)
        sections.extend(_extract_form_block(grid, block_cols, remaining))
        used = sum(len(s["items"]) for s in sections)
        remaining = max_items - used
    return sections


def _formula_audit(formulas):
    if not formulas:
        return {"count": 0, "patterns": [], "possible_inconsistencies": []}
    by_col = defaultdict(Counter)
    for f in formulas:
        by_col[f["column"]][f["pattern"]] += 1

    anomalies = []
    for col, counter in by_col.items():
        if sum(counter.values()) < 4:
            continue
        dominant, dom_count = counter.most_common(1)[0]
        if dom_count < 3:
            continue
        for f in formulas:
            if f["column"] == col and f["pattern"] != dominant and counter[f["pattern"]] <= 2:
                anomalies.append({
                    "cell": f["cell"],
                    "formula": f["formula"],
                    "reason": "padrão diferente do dominante na mesma coluna",
                })
                if len(anomalies) >= 30:
                    break
        if len(anomalies) >= 30:
            break

    patterns = Counter(f["pattern"] for f in formulas).most_common(12)
    return {
        "count": len(formulas),
        "patterns": [{"pattern": p, "count": c} for p, c in patterns],
        "possible_inconsistencies": anomalies,
        "note": "Fórmulas são lidas e comparadas, mas não são recalculadas pelo NÚCLEO.",
    }


def analyze_xlsx(path: Path):
    with _safe_xlsx_zip(path) as z:
        shared = _xlsx_shared_strings(z)
        style_kinds = _xlsx_number_formats(z)
        sheets = []
        total_rows = 0
        total_formulas = 0
        structures = Counter()
        for sheet_name, target in _xlsx_sheet_map(z):
            grid, max_row, max_col, formulas = _xlsx_grid(z, target, shared, style_kinds)
            structure = _classify_structure(grid, max_row, max_col)
            audit = _formula_audit(formulas)
            if structure == "form":
                sections = _extract_form(grid, max_row, max_col)
                field_count = sum(
                    1 for s in sections for it in s["items"]
                    if not it.get("structural")
                )
                sheet = {
                    "kind": "form",
                    "structure": "form",
                    "name": sheet_name,
                    "row_count": max_row,
                    "section_count": len(sections),
                    "field_count": field_count,
                    "sections": sections,
                    "formula_audit": audit,
                }
                total_rows += max_row
            else:
                rows = _grid_to_rows(grid, max_row, max_col)
                table = _analyze_rows(rows, source_name=sheet_name)
                table["structure"] = "table"
                table["formula_audit"] = audit
                sheet = table
                total_rows += table["row_count"]
            structures[structure] += 1
            sheets.append(sheet)
            total_formulas += audit["count"]
        return {
            "type": "xlsx",
            "sheet_count": len(sheets),
            "row_count": total_rows,
            "formula_count": total_formulas,
            "structures": dict(structures),
            "sheets": sheets,
            "note": "Leitura local sem executar macros ou código. Fórmulas não são recalculadas.",
        }


def _json_depth(obj, depth=0):
    if depth >= MAX_JSON_DEPTH:
        return depth
    if isinstance(obj, dict) and obj:
        return max(_json_depth(v, depth + 1) for v in list(obj.values())[:50])
    if isinstance(obj, list) and obj:
        return max(_json_depth(v, depth + 1) for v in obj[:50])
    return depth


def analyze_json(path: Path):
    text, encoding = _decode_text(path.read_bytes())
    try:
        obj = json.loads(text)
    except json.JSONDecodeError as exc:
        raise FileAnalysisError(f"JSON inválido: {exc}") from exc

    if isinstance(obj, list) and obj and all(isinstance(x, dict) for x in obj[:100]):
        keys = []
        seen = set()
        for row in obj[:1000]:
            for k in row.keys():
                if str(k) not in seen:
                    seen.add(str(k))
                    keys.append(str(k))
                    if len(keys) >= MAX_COLS:
                        break
            if len(keys) >= MAX_COLS:
                break
        rows = [keys]
        for row in obj[:MAX_ROWS]:
            rows.append([row.get(k) for k in keys])
        table = _analyze_rows(rows, source_name=path.name)
        table.update({"type": "json", "encoding": encoding,
                      "json_shape": "array_of_objects", "structure": "table"})
        return table

    preview = json.dumps(obj, ensure_ascii=False, indent=2)[:6000]
    keys = list(obj.keys())[:100] if isinstance(obj, dict) else []
    return {
        "type": "json",
        "json_shape": type(obj).__name__,
        "encoding": encoding,
        "depth_approx": _json_depth(obj),
        "top_keys": [str(k) for k in keys],
        "preview": preview,
    }


def analyze_text(path: Path, kind: str):
    text, encoding = _decode_text(path.read_bytes())
    lines = text.splitlines()
    headings = []
    if kind == "md":
        headings = [ln.strip()[:180] for ln in lines if ln.lstrip().startswith("#")][:40]
    return {
        "type": kind,
        "encoding": encoding,
        "line_count": len(lines),
        "char_count": len(text),
        "word_count": len(re.findall(r"\w+", text, flags=re.UNICODE)),
        "headings": headings,
        "preview": text[:7000],
    }


def analyze_pdf(path: Path):
    """Extrai texto de PDF NATIVO (digital) via pypdf — Python puro, sem OCR.
    PDFs escaneados (só imagem) não têm texto extraível: retornam páginas vazias
    e um aviso claro (OCR está fora do escopo, por portabilidade)."""
    try:
        from pypdf import PdfReader
        from pypdf.errors import PdfReadError
    except Exception:
        raise FileAnalysisError(
            "Leitura de PDF indisponível: dependência 'pypdf' ausente no runtime. "
            "Instale com: runtime\\python\\python.exe -m pip install pypdf"
        )

    try:
        reader = PdfReader(str(path))
    except PdfReadError as exc:
        raise FileAnalysisError(f"PDF inválido ou corrompido: {exc}")
    except Exception as exc:  # noqa
        raise FileAnalysisError(f"Não foi possível abrir o PDF: {exc}")

    encrypted = bool(getattr(reader, "is_encrypted", False))
    if encrypted:
        try:
            # tenta senha vazia (muitos PDFs só têm restrição, não senha real)
            if reader.decrypt("") == 0:
                raise FileAnalysisError(
                    "PDF protegido por senha — não é possível extrair o texto."
                )
        except FileAnalysisError:
            raise
        except Exception:
            raise FileAnalysisError(
                "PDF protegido/criptografado — não é possível extrair o texto."
            )

    total_pages = len(reader.pages)
    n = min(total_pages, MAX_PDF_PAGES)
    partes, empty_pages, chars = [], [], 0
    processed = 0
    truncated = False
    for i in range(n):
        processed = i + 1
        try:
            t = reader.pages[i].extract_text() or ""
        except Exception:
            t = ""
        t = t.strip()
        if not t:
            empty_pages.append(i + 1)
            continue
        marca = f"[p.{i + 1}]\n{t}"
        partes.append(marca)
        chars += len(marca)
        if chars >= MAX_PDF_TEXT_CHARS:
            truncated = True
            break
    texto = "\n\n".join(partes)

    # heurística de escaneado: maioria das páginas sem texto
    scanned = total_pages > 0 and len(empty_pages) >= max(1, int(0.8 * total_pages))
    nota = ""
    if scanned:
        nota = ("PDF parece ESCANEADO (imagem sem texto extraível). Extração nativa "
                "não recupera o conteúdo — precisaria de OCR, fora do escopo.")
    elif truncated:
        nota = (f"PDF longo: texto truncado no teto (~{MAX_PDF_TEXT_CHARS//1000}k caracteres); "
                f"lidas {processed} de {total_pages} páginas.")
    elif total_pages > MAX_PDF_PAGES:
        nota = f"PDF grande: processadas {n} de {total_pages} páginas."

    # títulos prováveis: linhas curtas isoladas (heurística leve), sem ruído de rodapé
    def _titulo_valido(s: str) -> bool:
        low = s.lower()
        if "p a g e" in low or re.search(r"\bpage\b|\bpág(?:ina)?\b", low):
            return False
        if len(re.findall(r"[A-Za-zÀ-ÿ]{3,}", s)) < 2:  # precisa de >=2 palavras reais
            return False
        return s.isupper() or bool(re.match(r"^\d+(\.\d+)*\s+\S", s))
    headings = []
    for ln in texto.splitlines():
        s = ln.strip()
        if 6 <= len(s) <= 90 and not s.startswith("[p.") and _titulo_valido(s):
            headings.append(s[:120])
        if len(headings) >= 30:
            break

    return {
        "type": "pdf",
        "page_count": total_pages,
        "pages_with_text": len(partes),
        "empty_pages": empty_pages[:50],
        "encrypted": encrypted,
        "scanned_guess": scanned,
        "char_count": len(texto),
        "word_count": len(re.findall(r"\w+", texto, flags=re.UNICODE)),
        "headings": headings,
        "preview": texto[:12000],
        "note": nota,
    }


def analyze_docx(path: Path):
    """Extrai texto de .docx (Word OOXML) via python-docx: parágrafos, títulos
    (por estilo) e tabelas. Não executa macros (docx sem macro; .docm à parte)."""
    try:
        from docx import Document
        from docx.opc.exceptions import PackageNotFoundError
    except Exception:
        raise FileAnalysisError(
            "Leitura de DOCX indisponível: dependência 'python-docx' ausente no runtime."
        )
    try:
        doc = Document(str(path))
    except PackageNotFoundError:
        raise FileAnalysisError(
            "Não é um .docx válido (OOXML). Se for um .doc antigo (Word 97-2003), "
            "salve como .docx no Word/LibreOffice e reenvie."
        )
    except Exception as exc:  # noqa
        raise FileAnalysisError(f"Não foi possível abrir o .docx: {exc}")

    paras, headings = [], []
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if not t:
            continue
        style = (getattr(p.style, "name", "") or "").lower()
        if "heading" in style or "título" in style or "titulo" in style or style == "title":
            headings.append(t[:120])
        paras.append(t)

    table_rows, n_tables = [], 0
    for tbl in doc.tables:
        n_tables += 1
        for row in tbl.rows:
            cells = [(c.text or "").strip() for c in row.cells]
            if any(cells):
                table_rows.append(" | ".join(cells))

    body = "\n".join(paras)
    if table_rows:
        body += "\n\n[TABELAS]\n" + "\n".join(table_rows[:400])
    body = body[:MAX_PDF_TEXT_CHARS]

    props = getattr(doc, "core_properties", None)
    titulo = (getattr(props, "title", "") or "") if props else ""
    autor = (getattr(props, "author", "") or "") if props else ""

    return {
        "type": "docx",
        "paragraph_count": len(paras),
        "table_count": n_tables,
        "char_count": len(body),
        "word_count": len(re.findall(r"\w+", body, flags=re.UNICODE)),
        "headings": headings[:40],
        "title": titulo,
        "author": autor,
        "preview": body[:12000],
        "note": "",
    }


def analyze_file(path: Path):
    path = Path(path)
    if not path.exists() or not path.is_file():
        raise FileAnalysisError("Arquivo não encontrado.")
    size = path.stat().st_size
    if size > MAX_FILE_BYTES:
        raise FileAnalysisError("Arquivo excede o limite de 50 MB.")
    ext = path.suffix.lower()
    if ext == ".doc":
        raise FileAnalysisError(
            "Formato .doc (Word 97-2003) não é suportado. Abra no Word ou LibreOffice "
            "e use 'Salvar como' → .docx; depois reenvie."
        )
    if ext not in SUPPORTED_EXTENSIONS:
        raise FileAnalysisError(
            "Formato não suportado nesta versão. Use XLSX, CSV, TXT, MD, JSON, PDF ou DOCX."
        )
    if ext == ".xlsx":
        return analyze_xlsx(path)
    if ext == ".csv":
        return analyze_csv(path)
    if ext == ".json":
        return analyze_json(path)
    if ext == ".pdf":
        return analyze_pdf(path)
    if ext == ".docx":
        return analyze_docx(path)
    return analyze_text(path, ext.lstrip("."))


def _question_tokens(question: str):
    stop = {
        "para","como","qual","quais","que","dos","das","uma","uns","umas","com",
        "sem","por","de","da","do","em","no","na","nos","nas","e","ou","o","a",
        "os","as","um","ao","aos","me","mostre","analise","analisar",
    }
    words = re.findall(r"[\wÀ-ÿ]+", str(question or "").lower())
    return [w for w in words if len(w) >= 3 and w not in stop][:40]


def _score_name(name, tokens):
    s = str(name or "").lower()
    return sum(3 for t in tokens if t in s)


def _compact_table_context(table, question, max_chars=7000):
    tokens = _question_tokens(question)
    cols = table.get("columns") or []
    ranked = sorted(
        cols,
        key=lambda c: (_score_name(c.get("name"), tokens), 1 if c.get("numeric") else 0),
        reverse=True,
    )
    selected = ranked[:10]
    lines = [
        f"Fonte tabular: {table.get('name','dados')}",
        f"Linhas analisadas: {table.get('row_count',0)}"
        + (" (limite atingido)" if table.get("truncated") else ""),
        "Colunas relevantes:",
    ]
    for c in selected:
        if c.get("numeric"):
            lines.append(
                f"- {c['name']}: numérica; n={c.get('numeric_count')}; "
                f"soma={c.get('sum')}; média={c.get('mean')}; "
                f"mín={c.get('min')}; máx={c.get('max')}; faltantes={c.get('missing')}"
            )
        else:
            tops = ", ".join(
                f"{x.get('value')} ({x.get('count')})"
                for x in (c.get("top_values") or [])[:6]
            )
            lines.append(
                f"- {c['name']}: categórica/texto; faltantes={c.get('missing')}; "
                f"valores frequentes: {tops or '—'}"
            )

    audit = table.get("formula_audit")
    if audit and audit.get("count"):
        lines.append(f"Fórmulas: {audit.get('count')}.")
        anomalies = audit.get("possible_inconsistencies") or []
        if anomalies:
            lines.append("Possíveis inconsistências de padrão de fórmula:")
            for a in anomalies[:8]:
                lines.append(f"- {a.get('cell')}: ={a.get('formula')} ({a.get('reason')})")
        lines.append(audit.get("note", ""))

    samples = table.get("sample_rows") or []
    if samples:
        lines.append("Amostra de linhas:")
        for row in samples[:5]:
            compact = "; ".join(f"{k}={v}" for k, v in list(row.items())[:8])
            lines.append("- " + compact[:800])

    return "\n".join(lines)[:max_chars]


def _compact_form_context(sheet, question, max_chars=3200):
    """Contexto semântico para abas tipo formulário/calculadora (V0.6.5.2)."""
    lines = [
        f"Planilha (formulário/calculadora): {sheet.get('name','planilha')}",
        f"Seções: {sheet.get('section_count',0)} | "
        f"campos rótulo→valor: {sheet.get('field_count',0)}",
    ]
    for sec in sheet.get("sections", []):
        title = sec.get("title")
        if title:
            lines.append("")
            lines.append(title.upper() if not title.isupper() else title)
        for it in sec.get("items", []):
            lbl = it.get("label") or "(sem rótulo)"
            cell = it.get("cell")
            formula = it.get("formula")
            disp = it.get("formatted")
            if disp is None and it.get("value") is not None:
                disp = str(it["value"])
            if disp is None:
                # rótulo/estrutural sem valor associado
                if formula:
                    lines.append(f"- {lbl} [{cell}]: fórmula ={formula}")
                elif not it.get("structural"):
                    lines.append(f"{lbl}")
                # estrutural puro (sem valor) é omitido para não virar "faltante"
            else:
                extra = f" (fórmula ={formula})" if formula else ""
                lines.append(f"- {lbl}: {disp} [{cell}]{extra}")

    audit = sheet.get("formula_audit")
    if audit and audit.get("count"):
        anomalies = audit.get("possible_inconsistencies") or []
        if anomalies:
            lines.append("")
            lines.append("Possíveis inconsistências de fórmula:")
            for a in anomalies[:6]:
                lines.append(f"- {a.get('cell')}: ={a.get('formula')}")

    lines.append("")
    lines.append(
        "Observação: rótulos, valores e fórmulas foram lidos do arquivo. "
        "O NÚCLEO calcula estatísticas de forma determinística, mas não "
        "recalcula o motor de fórmulas do Excel. Células vazias intencionais "
        "(espaçadores/estrutura) não são dados faltantes."
    )
    return "\n".join(lines)[:max_chars]


def context_for_analysis(analysis: dict, filename: str, question: str, max_chars=8500):
    typ = analysis.get("type")
    header = f"ARQUIVO: {filename} | tipo={typ}"
    if typ == "xlsx":
        sheets = analysis.get("sheets") or []
        tokens = _question_tokens(question)
        ranked = sorted(
            sheets,
            key=lambda s: _score_name(s.get("name"), tokens),
            reverse=True,
        )
        structures = analysis.get("structures") or {}
        struct_desc = ", ".join(f"{k}:{v}" for k, v in structures.items()) or "—"
        chunks = [
            header,
            f"Planilhas: {analysis.get('sheet_count',0)} | "
            f"linhas analisadas: {analysis.get('row_count',0)} | "
            f"fórmulas: {analysis.get('formula_count',0)} | "
            f"estrutura: {struct_desc}",
            analysis.get("note", ""),
        ]
        for sheet in ranked[:3]:
            if sheet.get("structure") == "form":
                chunks.append(_compact_form_context(sheet, question, max_chars=3400))
            else:
                chunks.append(_compact_table_context(sheet, question, max_chars=3200))
        return "\n\n".join(chunks)[:max_chars]

    if analysis.get("kind") == "table":
        return (header + "\n" + _compact_table_context(analysis, question, max_chars=max_chars-200))[:max_chars]

    if typ in ("txt", "md", "pdf", "docx"):
        preview = analysis.get("preview", "")
        tokens = _question_tokens(question)
        if tokens:
            lines = preview.splitlines()
            scored = sorted(
                ((sum(1 for t in tokens if t in ln.lower()), i, ln) for i, ln in enumerate(lines)),
                reverse=True,
            )
            hits = [ln for score, _, ln in scored if score > 0][:20]
        else:
            hits = []
        if typ == "pdf":
            meta = (f"Páginas={analysis.get('page_count')} "
                    f"(com texto={analysis.get('pages_with_text')}) | "
                    f"palavras={analysis.get('word_count')}")
        elif typ == "docx":
            meta = (f"Parágrafos={analysis.get('paragraph_count')} | "
                    f"tabelas={analysis.get('table_count')} | "
                    f"palavras={analysis.get('word_count')}")
        else:
            meta = f"Linhas={analysis.get('line_count')} | palavras={analysis.get('word_count')}"
        parts = [header, meta]
        if analysis.get("note"):
            parts.append("⚠ " + analysis["note"])
        if analysis.get("headings"):
            parts.append("Títulos: " + " | ".join(analysis["headings"][:12]))
        parts.append("Trechos relevantes:\n" + ("\n".join(hits) if hits else preview[:6000]))
        return "\n".join(parts)[:max_chars]

    if typ == "json":
        parts = [
            header,
            f"Estrutura={analysis.get('json_shape')} | profundidade~{analysis.get('depth_approx','—')}",
        ]
        if analysis.get("top_keys"):
            parts.append("Chaves: " + ", ".join(analysis["top_keys"][:40]))
        parts.append("Prévia:\n" + str(analysis.get("preview", ""))[:6000])
        return "\n".join(parts)[:max_chars]

    return header + "\nSem contexto analítico disponível."
