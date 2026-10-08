#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
exporters.py — Exportação de conteúdo (Markdown/texto) para .docx e .xlsx.

Usado pela rota POST /api/export do NÚCLEO para o Chat e a Análise de Arquivos.
DOCX: python-docx (já no runtime). XLSX: openpyxl.
Entrada: texto Markdown (respostas do chat / laudo da análise).
Saída: bytes do arquivo (para download).

Semântica:
- DOCX: prosa/relatório — títulos (#), listas, negrito/itálico/código inline,
  blocos de código e TABELAS markdown viram tabelas do Word.
- XLSX: foco tabular — cada tabela markdown vira uma aba; se não houver tabela,
  o texto entra numa aba única, uma linha por linha (para não sair vazio).
"""
from __future__ import annotations

import io
import re
from typing import List

# ---------------------------------------------------------------------------
# Parsing mínimo de Markdown (linha a linha) — suficiente p/ saída de LLM.
# ---------------------------------------------------------------------------

_H_RE = re.compile(r"^(#{1,6})\s+(.*)$")
_BULLET_RE = re.compile(r"^\s*[-*+]\s+(.*)$")
_NUM_RE = re.compile(r"^\s*\d+[.)]\s+(.*)$")
_FENCE_RE = re.compile(r"^\s*```")
_INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|__[^_]+__|\*[^*]+\*|_[^_]+_|`[^`]+`)")


def _is_table_sep(line: str) -> bool:
    """True se a linha é o separador de cabeçalho de tabela markdown: |---|:--:|"""
    s = line.strip()
    if "|" not in s or "-" not in s:
        return False
    core = s.strip("|")
    cells = [c.strip() for c in core.split("|")]
    if not cells:
        return False
    return all(c and set(c) <= set("-: ") and "-" in c for c in cells)


def _split_row(line: str) -> List[str]:
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    # não quebra em \| escapado
    parts = re.split(r"(?<!\\)\|", s)
    return [p.strip().replace("\\|", "|") for p in parts]


def parse_md_tables(md: str) -> List[List[List[str]]]:
    """Extrai todas as tabelas markdown. Retorna lista de tabelas;
    cada tabela = lista de linhas; cada linha = lista de células (strings)."""
    lines = md.splitlines()
    tables: List[List[List[str]]] = []
    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]
        if "|" in line and i + 1 < n and _is_table_sep(lines[i + 1]):
            header = _split_row(line)
            rows = [header]
            j = i + 2
            while j < n and "|" in lines[j] and lines[j].strip():
                rows.append(_split_row(lines[j]))
                j += 1
            # normaliza nº de colunas
            width = max(len(r) for r in rows)
            rows = [r + [""] * (width - len(r)) for r in rows]
            tables.append(rows)
            i = j
        else:
            i += 1
    return tables


def _strip_md(s: str) -> str:
    """Remove ênfase inline de markdown (**bold**, *it*, __x__, `code`) do texto da célula."""
    s = str(s or "")
    s = re.sub(r"\*\*([^*]+)\*\*", r"\1", s)
    s = re.sub(r"__([^_]+)__", r"\1", s)
    s = re.sub(r"\*([^*]+)\*", r"\1", s)
    s = re.sub(r"`([^`]+)`", r"\1", s)
    return s.strip()


def _num(v: str):
    """Converte string numérica/monetária em int/float quando faz sentido; senão devolve o
    texto limpo (sem markdown). Ex.: 'R$ 5.000,00' -> 5000.0 ; '**Receitas**' -> 'Receitas'."""
    s = (v or "").strip()
    if not s:
        return ""
    # tira moeda/percent/espaços para tentar interpretar como número
    t = re.sub(r"[R$€£%\s]", "", s)
    if not t or not re.search(r"\d", t):
        return _strip_md(s)
    # remove separador de milhar e normaliza vírgula decimal (pt-BR)
    if re.fullmatch(r"-?\d{1,3}(\.\d{3})+(,\d+)?", t):
        t = t.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d+,\d+", t):
        t = t.replace(",", ".")
    try:
        if re.fullmatch(r"-?\d+", t):
            return int(t)
        f = float(t)
        return f
    except ValueError:
        return _strip_md(s)


# ---------------------------------------------------------------------------
# DOCX (python-docx)
# ---------------------------------------------------------------------------

def _docx_runs(paragraph, text: str) -> None:
    """Adiciona 'text' a um parágrafo interpretando negrito/itálico/código inline."""
    pos = 0
    for m in _INLINE_RE.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        tok = m.group(0)
        if tok.startswith("**") or tok.startswith("__"):
            paragraph.add_run(tok[2:-2]).bold = True
        elif tok.startswith("`"):
            r = paragraph.add_run(tok[1:-1]); r.font.name = "Consolas"
        else:  # * _ itálico
            paragraph.add_run(tok[1:-1]).italic = True
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def md_to_docx(md: str, titulo: str = "NÚCLEO IA PORTÁTIL") -> bytes:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    lines = (md or "").replace("\r\n", "\n").replace("\r", "\n").split("\n")
    i, n = 0, len(lines)
    while i < n:
        line = lines[i]

        # bloco de código ```
        if _FENCE_RE.match(line):
            i += 1
            buf = []
            while i < n and not _FENCE_RE.match(lines[i]):
                buf.append(lines[i]); i += 1
            i += 1  # pula a fence de fechamento
            p = doc.add_paragraph()
            r = p.add_run("\n".join(buf))
            r.font.name = "Consolas"; r.font.size = Pt(9)
            continue

        # tabela markdown
        if "|" in line and i + 1 < n and _is_table_sep(lines[i + 1]):
            header = _split_row(line)
            rows = [header]
            j = i + 2
            while j < n and "|" in lines[j] and lines[j].strip():
                rows.append(_split_row(lines[j])); j += 1
            width = max(len(r) for r in rows)
            rows = [r + [""] * (width - len(r)) for r in rows]
            table = doc.add_table(rows=len(rows), cols=width)
            try:
                table.style = "Light Grid Accent 1"
            except Exception:
                pass
            for ri, row in enumerate(rows):
                for ci, cell in enumerate(row):
                    c = table.cell(ri, ci)
                    c.text = ""
                    para = c.paragraphs[0]
                    _docx_runs(para, cell)
                    if ri == 0:
                        for rn in para.runs:
                            rn.bold = True
            i = j
            doc.add_paragraph()
            continue

        m = _H_RE.match(line)
        if m:
            doc.add_heading(m.group(2).strip(), level=min(len(m.group(1)), 4))
            i += 1
            continue

        mb = _BULLET_RE.match(line)
        if mb:
            _docx_runs(doc.add_paragraph(style="List Bullet"), mb.group(1))
            i += 1
            continue

        mn = _NUM_RE.match(line)
        if mn:
            _docx_runs(doc.add_paragraph(style="List Number"), mn.group(1))
            i += 1
            continue

        if not line.strip():
            i += 1
            continue

        _docx_runs(doc.add_paragraph(), line)
        i += 1

    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()


# ---------------------------------------------------------------------------
# XLSX (openpyxl)
# ---------------------------------------------------------------------------

def _sheet_title(base: str, used: set) -> str:
    # Excel: máx 31 chars, sem : \ / ? * [ ]
    t = re.sub(r'[:\\/?*\[\]]', " ", base).strip()[:31] or "Planilha"
    orig, k = t, 2
    while t.lower() in used:
        suf = f" {k}"
        t = (orig[:31 - len(suf)] + suf)
        k += 1
    used.add(t.lower())
    return t


def md_to_xlsx(md: str, titulo: str = "NÚCLEO IA PORTÁTIL") -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment

    wb = Workbook()
    wb.remove(wb.active)
    used = set()
    tables = parse_md_tables(md or "")

    if tables:
        for idx, rows in enumerate(tables, 1):
            ws = wb.create_sheet(_sheet_title(f"Tabela {idx}", used))
            for ri, row in enumerate(rows):
                for ci, cell in enumerate(row, 1):
                    val = _strip_md(cell) if ri == 0 else _num(cell)
                    c = ws.cell(row=ri + 1, column=ci, value=val)
                    if ri == 0:
                        c.font = Font(bold=True)
                        c.alignment = Alignment(vertical="center")
            ws.freeze_panes = "A2"
            # largura aproximada por coluna
            for ci in range(1, len(rows[0]) + 1):
                width = max((len(str(r[ci - 1])) for r in rows if ci - 1 < len(r)), default=8)
                ws.column_dimensions[ws.cell(row=1, column=ci).column_letter].width = min(max(width + 2, 10), 60)
    else:
        # sem tabela: despeja o texto (uma linha por linha) para não sair vazio
        ws = wb.create_sheet(_sheet_title("Conteúdo", used))
        ws.column_dimensions["A"].width = 100
        for ln in (md or "").splitlines() or [""]:
            ws.append([ln])

    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()
